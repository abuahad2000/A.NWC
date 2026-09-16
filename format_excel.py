import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd

excel_path = r"c:/antigravity files IDE/PROGRAM/XLSX/بلاغات_المشاريع_الجارية_المعتمدة_المحدثة.xlsx"
df = pd.read_excel(excel_path)

# Create workbook with openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "البلاغات المعتمدة للمشاريع"
ws.views.sheetView[0].rightToLeft = True  # RTL layout

# Fonts
font_title = Font(name="Sakkal Majalla", size=18, bold=True, color="1B365D")
font_subtitle = Font(name="Sakkal Majalla", size=13, italic=True, color="495057")
font_header = Font(name="Sakkal Majalla", size=14, bold=True, color="FFFFFF")
font_data = Font(name="Sakkal Majalla", size=12, bold=False, color="000000")
font_data_bold = Font(name="Sakkal Majalla", size=12, bold=True, color="000000")
font_badge_amber = Font(name="Sakkal Majalla", size=12, bold=True, color="7A4100")
font_badge_blue = Font(name="Sakkal Majalla", size=12, bold=True, color="0C5460")
font_action_red = Font(name="Sakkal Majalla", size=12, bold=True, color="721C24")

# Fills
fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
fill_zebra = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
fill_amber_badge = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
fill_blue_badge = PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid")
fill_red_badge = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
fill_green_badge = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")

# Borders
thin_gray = Side(style='thin', color='D9D9D9')
border_cell = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

# Alignments
align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center')

# Title block (Rows 1 & 2)
ws.merge_cells("A1:K1")
ws["A1"] = "سجل بلاغات التعدي المعتمدة لقطاع المشاريع الجارية (NWC)"
ws["A1"].font = font_title
ws["A1"].alignment = Alignment(horizontal='right', vertical='center')

ws.merge_cells("A2:K2")
ws["A2"] = f"تاريخ الاستخراج: 16 سبتمبر 2026 | الحالات: (تحت معالجة المقاول - بانتظار اعتماد الجهة المتعدية) | المشاريع الجارية المعتمدة مكانياً"
ws["A2"].font = font_subtitle
ws["A2"].alignment = Alignment(horizontal='right', vertical='center')

ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 22
ws.row_dimensions[3].height = 10  # spacing row

# Write Table Headers at Row 4
headers = df.columns.tolist()
header_row = 4
for col_num, h_text in enumerate(headers, 1):
    cell = ws.cell(row=header_row, column=col_num)
    cell.value = str(h_text)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = border_cell

ws.row_dimensions[header_row].height = 32

# Write Data Rows
start_row = 5
for r_idx, row in df.iterrows():
    curr_row = start_row + r_idx
    ws.row_dimensions[curr_row].height = 26
    is_even = (r_idx % 2 == 0)
    row_fill = fill_white if is_even else fill_zebra
    
    for c_idx, val in enumerate(row, 1):
        cell = ws.cell(row=curr_row, column=c_idx)
        cell.value = val if pd.notnull(val) else "-"
        cell.font = font_data
        cell.fill = row_fill
        cell.border = border_cell
        
        col_name = headers[c_idx - 1]
        
        # Center align IDs, numbers, dates
        if col_name in ['id', 'operation_number', 'po', 'رقم_بلاغ_التعدي', 'تاريخ_البلاغ', 'status', 'خط_الطول', 'خط_العرض', 'جهة_التوجيه_والتنبيه', 'حالة_البلاغ']:
            cell.alignment = align_center
        else:
            cell.alignment = align_right
            
        # Status Badges
        if col_name == 'حالة_البلاغ':
            if val == 'تحت معالجة المقاول':
                cell.fill = fill_amber_badge
                cell.font = font_badge_amber
            elif val == 'بانتظار اعتماد الجهة المتعدية':
                cell.fill = fill_blue_badge
                cell.font = font_badge_blue
                
        # Action Target Badges
        if col_name == 'جهة_التوجيه_والتنبيه':
            if 'مدير البرنامج' in str(val):
                cell.fill = fill_amber_badge
                cell.font = font_badge_amber
            elif 'مسؤول التعديات' in str(val):
                cell.fill = fill_blue_badge
                cell.font = font_badge_blue

        if col_name == 'id' or col_name == 'رقم_بلاغ_التعدي':
            cell.font = font_data_bold

end_row = start_row + len(df) - 1
last_col_letter = get_column_letter(len(headers))

# Add Excel Dynamic Table
tab_range = f"A{header_row}:{last_col_letter}{end_row}"
table = Table(displayName="OngoingEncroachmentsTable", ref=tab_range)

# Clean and modern table style
style = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False
)
table.tableStyleInfo = style
ws.add_table(table)

# Auto-adjust column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col[header_row-1:]:
        val_str = str(cell.value or '')
        # approximate Arabic character length
        length = len(val_str.encode('utf-8')) // 2 + 3
        if length > max_len:
            max_len = length
    ws.column_dimensions[col_letter].width = max(max_len, 14)

# Specific adjustments for key columns
for col_idx, col_name in enumerate(headers, 1):
    c_letter = get_column_letter(col_idx)
    if 'name' in col_name or 'المشروع' in col_name:
        ws.column_dimensions[c_letter].width = 38
    elif 'وصف' in col_name or 'الإجراء' in col_name:
        ws.column_dimensions[c_letter].width = 45
    elif 'location' in col_name or 'الشارع' in col_name or 'الحي' in col_name:
        ws.column_dimensions[c_letter].width = 22
    elif 'contractor' in col_name or 'consultant' in col_name:
        ws.column_dimensions[c_letter].width = 30
    elif 'id' in col_name or 'po' in col_name or 'رقم' in col_name:
        ws.column_dimensions[c_letter].width = 16

# Save formatted file
wb.save(excel_path)
print(f"File successfully formatted with Sakkal Majalla and Dynamic Table at:\n{excel_path}")
